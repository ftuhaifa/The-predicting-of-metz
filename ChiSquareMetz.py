# -*- coding: utf-8 -*-
"""
Created on Sun Aug  6 14:39:13 2023

@author: ftuha
"""

import pandas as pd

# Organize the chi-square test results into a dictionary
results = {
    'Metastasis': ['DX-brain', 'DX-brain', 'DX-brain', 'DX-bone', 'DX-bone', 'DX-bone', 'DX-liver', 'DX-liver', 'DX-liver'],
    'Categorical Variable': ['Age', 'Sex', 'Race', 'Age', 'Sex', 'Race', 'Age', 'Sex', 'Race'],
    'Chi-square': [737.7369223849935, 67.02488337198272, 289.23889744511985, 212.30511895365433, 403.5793469419112, 339.71032817819594, 89.43839660521002, 353.90805153027145, 99.1467155045513],
    'P-value': [3.660580229438283e-150, 2.6810165762410666e-16, 2.268771857875985e-61, 9.327085719141967e-39, 9.157235442414892e-90, 2.920570022586747e-72, 6.341302750780829e-14, 5.971993279106593e-79, 1.494466098934493e-20],
    '95% Confidence Interval': [(716.7108525675104, 758.7629922024765), (63.1834245512886, 70.86634219267685), (279.7511684083387, 298.726626481901), (191.27904913617127, 233.3311887711374), (399.73788812121705, 407.4208057626053), (330.2225991414148, 349.1980572149771), (68.41232678772695, 110.46446642269308), (350.0665927095773, 357.7495103509656), (89.65898646777015, 108.63444454133244)]
}

# Create a pandas DataFrame from the results dictionary with a MultiIndex
df_results = pd.DataFrame(results)
df_results = df_results.set_index(['Categorical Variable', 'Metastasis'])

# Reshape the DataFrame to have each metastasis as a separate column
df_results_pivoted = df_results.unstack(level=1)

# Display the DataFrame as a table
print(df_results_pivoted)


from docx import Document
import pandas as pd

# ... (The code to create the DataFrame remains the same)

# Create a new Word document
doc = Document()

# Convert the DataFrame to a table and add it to the Word document
table = doc.add_table(rows=len(df_results_pivoted)+1, cols=len(df_results_pivoted.columns), style='Table Grid')
table.cell(0, 0).text = "Categorical Variable"
for i, (metastasis, group) in enumerate(df_results_pivoted.groupby(level=1, axis=1)):
    table.cell(0, i+1).text = metastasis
    for j, col in enumerate(group.columns):
        table.cell(j+1, 0).text = col
        for k, v in group[col].iteritems():
            table.cell(j+1, i+1).text = str(v)

# Save the Word document
doc.save("ChiSquareResults22.docx")
